"""Script pour générer la présentation PowerPoint de soutenance."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def add_title_slide(prs, title, subtitle):
    """Ajoute une diapositive de titre."""
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_shape = slide.shapes.title
    subtitle_shape = slide.placeholders[1]
    
    title_shape.text = title
    subtitle_shape.text = subtitle
    
    title_shape.text_frame.paragraphs[0].font.size = Pt(54)
    title_shape.text_frame.paragraphs[0].font.bold = True
    title_shape.text_frame.paragraphs[0].font.color.rgb = RGBColor(31, 78, 121)
    
    return slide

def add_bullet_slide(prs, title, bullets):
    """Ajoute une diapositive avec titre et points de liste."""
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title_shape = slide.shapes.title
    body_shape = slide.placeholders[1]
    
    title_shape.text = title
    title_shape.text_frame.paragraphs[0].font.size = Pt(44)
    title_shape.text_frame.paragraphs[0].font.bold = True
    title_shape.text_frame.paragraphs[0].font.color.rgb = RGBColor(31, 78, 121)
    
    text_frame = body_shape.text_frame
    text_frame.clear()
    
    for i, bullet in enumerate(bullets):
        if i == 0:
            p = text_frame.paragraphs[0]
        else:
            p = text_frame.add_paragraph()
        p.text = bullet
        p.level = 0
        p.font.size = Pt(24)
        p.space_before = Pt(12)
    
    return slide

def add_section_slide(prs, title):
    """Ajoute une diapositive de section."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(31, 78, 121)
    
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(9), Inches(1.5))
    title_frame = title_box.text_frame
    p = title_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(54)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.alignment = PP_ALIGN.CENTER
    
    return slide

def create_presentation():
    """Crée la présentation complète."""
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)
    
    # Diapositive 1: Titre
    add_title_slide(prs, 
        "Classification Automatique des Réclamations",
        "Système NLP/MLOps complet\nSoutenance de Stage - Data Science"
    )
    
    # Diapositive 2: Contexte et objectifs
    add_bullet_slide(prs, "Contexte et Objectifs",
        [
            "🎯 Automatiser le traitement des réclamations clients",
            "📊 Réduire le temps de traitement et améliorer la réactivité",
            "🔀 Routage intelligent vers les équipes compétentes",
            "⚡ Système de priorité hybride (règles + ML)",
            "🏗️ Architecture MLOps production-ready"
        ]
    )
    
    # Diapositive 3: Architecture globale
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "Architecture Globale"
    title.text_frame.paragraphs[0].font.size = Pt(44)
    title.text_frame.paragraphs[0].font.bold = True
    title.text_frame.paragraphs[0].font.color.rgb = RGBColor(31, 78, 121)
    
    body = slide.placeholders[1].text_frame
    body.clear()
    
    archi_text = """
┌─────────────────────────────────────────────────┐
│       Dashboard Streamlit (Port 8501)           │
│   • Interface utilisateur interactive           │
│   • Single et batch processing                  │
│   • Statistiques temps-réel                     │
└────────────────┬────────────────────────────────┘
                 │ (API_BASE_URL=http://api:8000)
┌────────────────▼────────────────────────────────┐
│     API FastAPI (Port 8000)                     │
│   • Endpoints: /classify, /batch, /stats        │
│   • Health check et monitoring                  │
│   • CORS enabled                                │
└────────────────┬────────────────────────────────┘
                 │
    ┌────────────┴────────────┐
    ▼                         ▼
┌─────────────┐      ┌──────────────────┐
│ Classifier  │      │ Entity Extractor │
│ (TF-IDF +   │      │ (spaCy + regex)  │
│ LogReg)     │      │ fallback         │
└─────────────┘      └──────────────────┘
"""
    
    p = body.paragraphs[0]
    p.text = archi_text
    p.font.name = 'Courier New'
    p.font.size = Pt(14)
    
    # Diapositive 4: Stack technique
    add_bullet_slide(prs, "Stack Technique",
        [
            "🐍 Python 3.11 (94.5% du code)",
            "⚙️ FastAPI (API REST optimisée) + Uvicorn",
            "📊 scikit-learn (classification : TF-IDF + Logistic Regression)",
            "🧠 spaCy (extraction d'entités NER - modèle custom entraîné)",
            "🎨 Streamlit (interface utilisateur interactive)",
            "🐳 Docker + Docker Compose (orchestration)"
        ]
    )
    
    # Diapositive 5: Pipeline de traitement
    add_bullet_slide(prs, "Pipeline de Traitement",
        [
            "1️⃣ Nettoyage du texte (text_cleaning.py)",
            "2️⃣ Classification (TF-IDF vectorization + LogReg)",
            "3️⃣ Extraction d'entités (spaCy NER ou regex fallback)",
            "4️⃣ Scoring de priorité (règles métier + signal modèle)",
            "5️⃣ Routage automatique vers équipe compétente",
            "6️⃣ Escalade manager si nécessaire"
        ]
    )
    
    # Section: Modèle de classification
    add_section_slide(prs, "Modèle de Classification")
    
    # Diapositive 6: Classification
    add_bullet_slide(prs, "Classification des Réclamations",
        [
            "📌 Architecture : Pipeline scikit-learn (TF-IDF → LogReg)",
            "📦 Fichier : baseline_tfidf_logreg_pipeline.joblib",
            "🏷️ 8 catégories prédéfinies :",
            "    • PERTE_VOL, ERREUR_DESTINATION, RETARD_LIVRAISON",
            "    • PRODUIT_ENDOMMAGE, RETOUR_REMBOURSEMENT",
            "    • FACTURATION_PAIEMENT, QUALITE_SERVICE, AUTRE",
            "📈 Sortie : Probabilités pour chaque classe + confiance"
        ]
    )
    
    # Diapositive 7: Extraction d'entités
    add_bullet_slide(prs, "Extraction d'Entités (NER)",
        [
            "🎯 Modèle principal : spaCy custom entraîné",
            "📂 Artefact : ner/model-best/",
            "🔍 Entités extraites :",
            "    • NUMERO_TRACKING (formats : TRK-2025-51462, etc.)",
            "    • VILLE (nécessite modèle spaCy)",
            "    • MONTANT (€, EUR, dinars tunisiens DT)",
            "    • DATE (formats numériques et textuels)",
            "⚠️ Fallback automatique : regex si modèle absent"
        ]
    )
    
    # Diapositive 8: Priorité hybride
    add_bullet_slide(prs, "Système de Priorité Hybride",
        [
            "🔄 Approche combinée : Règles métier (80%) + Modèle (20%)",
            "🔴 Déclencheurs de priorité HAUTE (règles dures) :",
            "    • Dommage majeur détecté (explosion, incendie, blessure)",
            "    • Retard critique (≥14 jours)",
            "    • Escalade explicite (avocat, presse, tribunaux)",
            "    • Perte/vol confirmé",
            "📊 Sinon : score continu basé sur signaux textuels + catégorie"
        ]
    )
    
    # Diapositive 9: Signaux de priorité
    add_bullet_slide(prs, "Signaux d'Analyse de Priorité",
        [
            "✅ Signaux textuels capturés :",
            "    • Dommages mineurs (casse, rayure, fuite)",
            "    • Retard modéré (3-14 jours)",
            "    • Montant déclaré (seuil : 300€)",
            "    • Ton agressif (ratio majuscules > 30% ou 2+ points d'exclamation)",
            "    • Erreur de destination",
            "✅ Signal modèle : sévérité de catégorie × confiance prédiction"
        ]
    )
    
    # Diapositive 10: Routage équipes
    add_bullet_slide(prs, "Routage Automatique vers Équipes",
        [
            "🏢 Table de routage catégorie → équipe :",
            "    • Perte/Vol, Erreur destination → Logistique",
            "    • Produit endommagé → Qualité / SAV",
            "    • Retour/Remboursement, Facturation → Finance",
            "    • Qualité de service → Service Après-Vente",
            "    • Autres → Support général",
            "⚠️ Escalade manager : dommage majeur + plainte forte"
        ]
    )
    
    # Section: Dashboard et Monitoring
    add_section_slide(prs, "Dashboard & Monitoring")
    
    # Diapositive 11: Dashboard Streamlit
    add_bullet_slide(prs, "Dashboard Streamlit",
        [
            "🔍 Onglet 1 : Analyser une réclamation unique",
            "    • Saisie manuelle + bouton exemple",
            "    • Affichage : catégorie, confiance, priorité, équipe, entités",
            "📦 Onglet 2 : Traitement en lot",
            "    • Source : saisie multi-ligne ou fichier CSV",
            "    • Résultats : tableau filtrable + téléchargement CSV",
            "📊 Onglet 3 : Statistiques globales",
            "    • Compteurs en mémoire (cumul depuis démarrage API)"
        ]
    )
    
    # Diapositive 12: Santé du service
    add_bullet_slide(prs, "Health Check & Monitoring",
        [
            "🏥 Endpoint /health :",
            "    • État du service : OK ou DÉGRADÉ",
            "    • Modèle de classification chargé : ✅ ou ❌",
            "    • Mode NER : spaCy complet ou regex fallback",
            "📊 Endpoint /api/stats :",
            "    • Total traité, confiance moyenne, escalades manager",
            "    • Distributions par catégorie/priorité/équipe",
            "    • Déclencheurs de règles dures"
        ]
    )
    
    # Section: Infrastructure & Déploiement
    add_section_slide(prs, "Infrastructure & Déploiement")
    
    # Diapositive 13: Docker & Déploiement
    add_bullet_slide(prs, "Docker & Déploiement",
        [
            "🐳 Docker Compose orchestration :",
            "    • Service API (8000) : FastAPI + modèles",
            "    • Service Dashboard (8501) : Streamlit",
            "    • Health checks intégrés",
            "📂 Structure :",
            "    • api/ : Code API + requirements.txt + Dockerfile",
            "    • dashboard/ : Code Streamlit + requirements.txt + Dockerfile",
            "    • docker-compose.yml : Orchestration multi-conteneur"
        ]
    )
    
    # Diapositive 14: Artefacts ML
    add_bullet_slide(prs, "Artefacts ML Requis",
        [
            "⚠️ Fichiers à placer dans ./api/models/ AVANT docker build :",
            "📦 baseline_tfidf_logreg_pipeline.joblib",
            "    • Pipeline complet (vectorizer TF-IDF + classificateur LogReg)",
            "    • Généré par notebook (section 10, section 15)",
            "📂 ner/model-best/",
            "    • Modèle spaCy entraîné sur domaine réclamations",
            "    • Généré par notebook section 16.5",
            "❌ Si absent : API en mode dégradé (fallback regex pour NER)"
        ]
    )
    
    # Section: Limitations & Perspectives
    add_section_slide(prs, "Limitations & Perspectives")
    
    # Diapositive 15: Limitations actuelles
    add_bullet_slide(prs, "Limitations Actuelles",
        [
            "❌ Pas d'API d'entraînement automatique (retraining)",
            "❌ Pas de persistance de données (stats en mémoire volatile)",
            "❌ Pas de cache ou optimisation de latence avancée",
            "❌ Pas de tests unitaires/d'intégration explicites",
            "❌ Pas de logging structuré ou agrégation centralisée",
            "❌ Pas de gestion d'A/B testing ou versioning de modèles",
            "❌ Pas de documentation API Swagger personnalisée (FastAPI génère basic)"
        ]
    )
    
    # Diapositive 16: Perspectives d'évolution
    add_bullet_slide(prs, "Perspectives d'Évolution",
        [
            "✨ Feedback en boucle : labellisation utilisateur → retraining",
            "✨ Persistance : intégration PostgreSQL/MongoDB pour historique",
            "✨ MLflow : suivi des expériences, versioning des modèles",
            "✨ Monitoring avancé : Prometheus + Grafana",
            "✨ Tests complets : pytest + CI/CD GitHub Actions",
            "✨ Load testing : simulation haute charge (>100 req/s)",
            "✨ Explainabilité : LIME/SHAP pour interprétabilité prédictions"
        ]
    )
    
    # Section: Résumé
    add_section_slide(prs, "Résumé")
    
    # Diapositive 17: Points clés
    add_bullet_slide(prs, "Points Clés du Projet",
        [
            "✅ Système NLP end-to-end : classification + NER + priorité + routage",
            "✅ Architecture MLOps modulaire : API + Dashboard + Orchestration",
            "✅ Robustesse : fallback automatique (regex si modèles absent)",
            "✅ Scalabilité : batch processing (jusqu'à 200 réclamations)",
            "✅ Maintenance : configuration centralisée, env vars",
            "✅ Production-ready : health checks, CORS, error handling"
        ]
    )
    
    # Diapositive 18: Conclusion
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "Conclusion"
    title.text_frame.paragraphs[0].font.size = Pt(44)
    title.text_frame.paragraphs[0].font.bold = True
    title.text_frame.paragraphs[0].font.color.rgb = RGBColor(31, 78, 121)
    
    body = slide.placeholders[1].text_frame
    body.clear()
    
    conclusion_points = [
        "🎯 Démonstration d'un pipeline ML complet en production",
        "📊 Approche hybride (ML + règles métier) pragmatique et maintenable",
        "🚀 Prêt pour un déploiement en production (avec améliorations)",
        "📈 Fondation solide pour l'IA appliquée au CRM"
    ]
    
    for i, point in enumerate(conclusion_points):
        if i == 0:
            p = body.paragraphs[0]
        else:
            p = body.add_paragraph()
        p.text = point
        p.font.size = Pt(24)
        p.space_before = Pt(12)
    
    # Sauvegarder
    prs.save('presentation_soutenance.pptx')
    print("✅ Présentation créée : presentation_soutenance.pptx")

if __name__ == "__main__":
    create_presentation()
